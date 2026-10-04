from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    MODEL_VERSION,
    DailyHistoryRow,
    DataIssue,
    HoldingAttributionRow,
    RunRecord,
    TradePnlRow,
)


COLUMNS = (
    ("portfolio_code", "组合代码"),
    ("trade_date", "归因日期"),
    ("previous_trade_date", "昨日交易日"),
    ("status", "状态"),
    ("holding_count", "昨日持仓明细数"),
    ("holding_option_count", "持仓期权数"),
    ("holding_etf_count", "持仓ETF数"),
    ("trade_count", "盘中流水数"),
    ("trade_option_count", "交易期权数"),
    ("trade_etf_count", "交易ETF数"),
    ("issue_count", "问题数"),
    ("greeks_unavailable_count", "Greeks不可用持仓数"),
    ("holding_pnl", "昨日持仓PnL"),
    ("holding_option_pnl", "期权持仓PnL"),
    ("holding_etf_pnl", "ETF持仓PnL（信息项）"),
    ("pnl_delta", "Delta PnL（含ETF）"),
    ("pnl_gamma", "Gamma PnL"),
    ("pnl_vega", "Vega PnL"),
    ("pnl_theta", "Theta PnL"),
    ("pnl_residual", "持仓Residual"),
    ("intraday_turnover", "盘中绝对成交额"),
    ("intraday_gross", "盘中交易毛收益（信息项）"),
    ("intraday_fee", "盘中交易费用"),
    ("intraday_net", "盘中交易净收益"),
    ("intraday_gross_return", "盘中毛收益率"),
    ("intraday_net_return", "盘中净收益率"),
    ("daily_gross", "每日总毛收益"),
    ("daily_net", "每日总净收益"),
    ("reconciliation", "公式勾稽差额"),
    ("updated_at", "更新时间"),
    ("model_version", "模型版本"),
)
HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MONEY_FORMAT = '#,##0.00;[Red](#,##0.00);-'
RATE_FORMAT = '0.0000%;[Red](0.0000%);-'
STATUS_RANK = {"FAILED": 0, "PARTIAL": 1, "SUCCESS": 2}
LEGACY_LABELS = {
    "holding_etf_pnl": "ETF持仓PnL",
    "pnl_delta": "Delta PnL",
    "intraday_gross": "盘中交易毛收益",
}


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _as_datetime(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    return datetime.fromisoformat(str(value))


def build_daily_rows(
    holding_rows: list[HoldingAttributionRow],
    trade_rows: list[TradePnlRow],
    issues: list[DataIssue],
    run_records: list[RunRecord],
) -> list[DailyHistoryRow]:
    holdings_by_date: dict[date, list[HoldingAttributionRow]] = defaultdict(list)
    trades_by_date: dict[date, list[TradePnlRow]] = defaultdict(list)
    issues_by_date: dict[date, list[DataIssue]] = defaultdict(list)
    for row in holding_rows:
        holdings_by_date[row.trade_date].append(row)
    for row in trade_rows:
        trades_by_date[row.trade_date].append(row)
    for issue in issues:
        issues_by_date[issue.trade_date].append(issue)

    result: list[DailyHistoryRow] = []
    for record in run_records:
        date_holdings = holdings_by_date[record.trade_date]
        date_trades = trades_by_date[record.trade_date]
        holding_pnl = sum(row.pnl_holding for row in date_holdings)
        holding_option_pnl = sum(
            row.pnl_holding
            for row in date_holdings
            if row.asset_type == "OPT_F"
        )
        holding_etf_pnl = sum(
            row.pnl_holding
            for row in date_holdings
            if row.asset_type == "SPT_ETF"
        )
        # ETF Delta=1：持仓ETF价格收益和盘中ETF毛收益都进入Delta。
        # 盘中期权仍按普通交易毛收益展示，不做日内Greeks归因。
        pnl_delta = sum(row.pnl_delta for row in date_holdings) + sum(
            row.pnl_delta for row in date_trades
        )
        pnl_gamma = sum(row.pnl_gamma for row in date_holdings)
        pnl_vega = sum(row.pnl_vega for row in date_holdings)
        pnl_theta = sum(row.pnl_theta for row in date_holdings)
        pnl_residual = sum(row.pnl_residual for row in date_holdings)
        intraday_turnover = sum(row.turnover for row in date_trades)
        intraday_gross = sum(row.pnl_gross for row in date_trades)
        option_intraday_gross = sum(
            row.pnl_gross
            for row in date_trades
            if row.asset_type == "OPT_F"
        )
        intraday_fee = sum(row.fee for row in date_trades)
        intraday_net = intraday_gross - intraday_fee
        gross_return = (
            intraday_gross / intraday_turnover
            if intraday_turnover
            else None
        )
        net_return = (
            intraday_net / intraday_turnover
            if intraday_turnover
            else None
        )
        daily_gross = holding_pnl + intraday_gross
        daily_net = daily_gross - intraday_fee
        explained_net = (
            pnl_delta
            + pnl_gamma
            + pnl_vega
            + pnl_theta
            + pnl_residual
            + option_intraday_gross
            - intraday_fee
        )
        result.append(
            DailyHistoryRow(
                portfolio_code=record.portfolio_code,
                trade_date=record.trade_date,
                previous_trade_date=record.previous_trade_date,
                status=record.status,
                holding_count=len(date_holdings),
                holding_option_count=sum(
                    row.asset_type == "OPT_F" for row in date_holdings
                ),
                holding_etf_count=sum(
                    row.asset_type == "SPT_ETF" for row in date_holdings
                ),
                trade_count=len(date_trades),
                trade_option_count=sum(
                    row.asset_type == "OPT_F" for row in date_trades
                ),
                trade_etf_count=sum(
                    row.asset_type == "SPT_S" for row in date_trades
                ),
                issue_count=len(issues_by_date[record.trade_date]),
                greeks_unavailable_count=sum(
                    row.asset_type == "OPT_F"
                    and row.iv_fit_status_start == "UNAVAILABLE"
                    for row in date_holdings
                ),
                holding_pnl=holding_pnl,
                holding_option_pnl=holding_option_pnl,
                holding_etf_pnl=holding_etf_pnl,
                pnl_delta=pnl_delta,
                pnl_gamma=pnl_gamma,
                pnl_vega=pnl_vega,
                pnl_theta=pnl_theta,
                pnl_residual=pnl_residual,
                intraday_turnover=intraday_turnover,
                intraday_gross=intraday_gross,
                intraday_fee=intraday_fee,
                intraday_net=intraday_net,
                intraday_gross_return=gross_return,
                intraday_net_return=net_return,
                daily_gross=daily_gross,
                daily_net=daily_net,
                reconciliation=daily_net - explained_net,
                updated_at=record.updated_at,
                model_version=record.model_version,
            )
        )
    return result


def load_daily_history(path: Path) -> dict[tuple[str, date], DailyHistoryRow]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        labels = tuple(next(values, ()))
        # v3 日汇总没有“Greeks不可用持仓数”。这里仍允许读取它，
        # 新模型重算后会按当前完整表头重新写回。
        old_columns = tuple(
            item for item in COLUMNS if item[0] != "greeks_unavailable_count"
        )

        def labels_match(columns) -> bool:
            return len(labels) == len(columns) and all(
                actual in {expected, LEGACY_LABELS.get(name, expected)}
                for actual, (name, expected) in zip(labels, columns)
            )

        if labels_match(COLUMNS):
            source_columns = COLUMNS
        elif labels_match(old_columns):
            source_columns = old_columns
        else:
            raise ValueError(
                f"{path} 是旧版或其他模型的日汇总文件；"
                "请使用新的 --daily-output 文件名"
            )
        names = [name for name, _ in source_columns]
        result: dict[tuple[str, date], DailyHistoryRow] = {}
        for values_row in values:
            if not any(value is not None for value in values_row):
                continue
            item = dict(zip(names, values_row))
            item.setdefault("greeks_unavailable_count", 0)
            item["portfolio_code"] = str(item["portfolio_code"])
            item["trade_date"] = _as_date(item["trade_date"])
            item["previous_trade_date"] = _as_date(
                item["previous_trade_date"]
            )
            item["updated_at"] = _as_datetime(item["updated_at"])
            for field in (
                "holding_count",
                "holding_option_count",
                "holding_etf_count",
                "trade_count",
                "trade_option_count",
                "trade_etf_count",
                "issue_count",
                "greeks_unavailable_count",
            ):
                item[field] = int(item[field] or 0)
            for field in (
                "holding_pnl",
                "holding_option_pnl",
                "holding_etf_pnl",
                "pnl_delta",
                "pnl_gamma",
                "pnl_vega",
                "pnl_theta",
                "pnl_residual",
                "intraday_turnover",
                "intraday_gross",
                "intraday_fee",
                "intraday_net",
                "daily_gross",
                "daily_net",
                "reconciliation",
            ):
                item[field] = float(item[field] or 0.0)
            for field in ("intraday_gross_return", "intraday_net_return"):
                item[field] = (
                    float(item[field]) if item[field] is not None else None
                )
            row = DailyHistoryRow(**item)
            result[(row.portfolio_code, row.trade_date)] = row
        return result
    finally:
        workbook.close()


def completed_dates(path: Path, portfolio: str) -> set[date]:
    return {
        item.trade_date
        for item in load_daily_history(path).values()
        if item.portfolio_code == portfolio
        and item.model_version == MODEL_VERSION
        and item.status == "SUCCESS"
    }


def update_daily_history(
    path: Path,
    holding_rows: list[HoldingAttributionRow],
    trade_rows: list[TradePnlRow],
    issues: list[DataIssue],
    run_records: list[RunRecord],
    force_dates: set[date] | None = None,
) -> int:
    history = load_daily_history(path)
    current = build_daily_rows(
        holding_rows, trade_rows, issues, run_records
    )
    updated_count = 0
    forced = force_dates or set()
    for row in current:
        key = (row.portfolio_code, row.trade_date)
        old = history.get(key)
        if (
            old is None
            or row.trade_date in forced
            or old.model_version != row.model_version
            or STATUS_RANK[row.status] >= STATUS_RANK[old.status]
        ):
            history[key] = row
            updated_count += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "每日增量汇总"
    worksheet.append([label for _, label in COLUMNS])
    for row in sorted(
        history.values(), key=lambda item: (item.trade_date, item.portfolio_code)
    ):
        values = asdict(row)
        worksheet.append([values[name] for name, _ in COLUMNS])

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    names = [name for name, _ in COLUMNS]
    date_columns = {
        names.index("trade_date") + 1,
        names.index("previous_trade_date") + 1,
    }
    updated_column = names.index("updated_at") + 1
    rate_columns = {
        names.index("intraday_gross_return") + 1,
        names.index("intraday_net_return") + 1,
    }
    money_columns = {
        index + 1
        for index, name in enumerate(names)
        if name.startswith("pnl_")
        or name.startswith("holding_")
        or name.startswith("intraday_")
        or name.startswith("daily_")
        or name == "reconciliation"
    } - rate_columns - {
        names.index("holding_count") + 1,
        names.index("holding_option_count") + 1,
        names.index("holding_etf_count") + 1,
        names.index("greeks_unavailable_count") + 1,
    }
    for cells in worksheet.iter_rows(min_row=2):
        for index in date_columns:
            cells[index - 1].number_format = "yyyy-mm-dd"
        cells[updated_column - 1].number_format = "yyyy-mm-dd hh:mm:ss"
        for index in money_columns:
            cells[index - 1].number_format = MONEY_FORMAT
        for index in rate_columns:
            cells[index - 1].number_format = RATE_FORMAT

    for column in worksheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        letter = get_column_letter(column[0].column)
        worksheet.column_dimensions[letter].width = min(max(width + 2, 10), 24)

    temporary = path.with_name(f".{path.stem}.tmp{path.suffix}")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    return updated_count



